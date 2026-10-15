import tempfile
import unittest
from pathlib import Path
from PIL import Image
from desktop.core import discover, prepare

class DesktopTest(unittest.TestCase):
    def test_order_crop_and_preservation(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);source=root/'source'/'ABC';source.mkdir(parents=True)
            for name in ['ABC-细节10.jpg','ABC-后.jpg','ABC-前1.jpg','ABC-前.jpg','ABC-细节2.jpg']:
                Image.new('RGB',(1350,1800),'red').save(source/name)
            before={p.name:p.read_bytes() for p in source.iterdir()}
            groups=discover(root/'source')
            self.assertEqual([p.name for p in groups[0]['files']],['ABC-前.jpg','ABC-前1.jpg','ABC-后.jpg','ABC-细节2.jpg','ABC-细节10.jpg'])
            ready=prepare(groups[0],root/'out',1000)
            with Image.open(ready['material']) as im:self.assertEqual(im.size,(1000,1000))
            self.assertEqual(before,{p.name:p.read_bytes() for p in source.iterdir()})
            self.assertEqual(len(ready['carousel']),5)
            with self.assertRaises(FileExistsError):prepare(groups[0],root/'out')

    def test_custom_main_suffix_drives_crop(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);source=root/'SKU';source.mkdir()
            for name,color in [('01.jpg','red'),('02-main.jpg','blue'),('03-main-detail.jpg','green')]:
                Image.new('RGB',(100,120),color).save(source/name)
            group=discover(source,'-MAIN')[0]
            self.assertEqual(group['files'][0].name,'02-main.jpg')
            ready=prepare(group,root/'output',800)
            with Image.open(ready['material']) as image:
                self.assertEqual(image.size,(800,800))
                self.assertGreater(image.getpixel((400,400))[2],240)
            for suffix in ('','missing'):
                self.assertEqual(discover(source,suffix)[0]['files'][0].name,'01.jpg')

    def test_bundled_suit_to_excel(self):
        from desktop.core import load_profile
        from import_tools.dianxiaomi import export
        from openpyxl import load_workbook
        assets=Path(__file__).resolve().parents[1]/'desktop'/'assets'
        config=load_profile(assets/'运动短袖短裤套装.xlsx')
        self.assertEqual([v['size'] for v in config['variants']],['S','M','L','XL','XXL'])
        for v in config['variants']:
            self.assertEqual(v['weight_g'],200)
            v['price']=39
        products=[{'sku':'YQ-TEST','print_name':'纽约字母','carousel_urls':'https://example.com/front.jpg','material_url':'https://example.com/material.jpg'}]
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'import.xlsx'
            self.assertEqual(export(assets/'店小秘创建商品模板.xlsx',config,products,out),5)
            w=load_workbook(out);s=w['导入模板'];headers=[c.value for c in s[1]]
            data=[dict(zip(headers,row)) for row in s.iter_rows(min_row=2,values_only=True)]
            self.assertEqual(data[0]['*产品标题'],'女士圆领短袖2件套纽约字母印花圆领T恤+短裤')
            self.assertTrue(all(r['*申报价格\n(店铺币种)']==39 for r in data))
            self.assertEqual(data[-1]['SKU货号'],'YQ-TEST-XXL')
            self.assertEqual(s.max_column,51)
            self.assertTrue(all(r['外包装形状']=='长方体' for r in data))
            self.assertTrue(all(r['外包装类型']=='软包装+软物' for r in data))
            self.assertTrue(all(r['外包装图片']=='https://img.cdnfe.com/product/open/ad7a1f90a1b744c08ed8a93b5c0fd982-goods.jpeg' for r in data))
