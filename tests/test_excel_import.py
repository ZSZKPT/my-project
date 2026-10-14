import tempfile
import unittest
from pathlib import Path
from import_tools.dianxiaomi import build_rows, export

class ExcelTest(unittest.TestCase):
    def setUp(self):
        self.config={'title':'女装{print_name} {sku}','variants':[{'size':s,'price':20,'length_cm':30,'width_cm':20,'height_cm':2,'weight_g':200} for s in ['S','M']]}
        self.products=[{'sku':'001','print_name':'爱心','carousel_urls':'https://example.com/1.jpg|https://example.com/2.jpg','material_url':'https://example.com/3.jpg'}]

    def test_mapping(self):
        rows=build_rows(self.config,self.products)
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[0]['产品货号'],'001')
        self.assertEqual(rows[1]['SKU货号'],'001-M')
        self.assertEqual(rows[0]['*产品标题'],rows[1]['*产品标题'])
        self.assertIn('\n',rows[0]['*轮播图'])

    def test_invalid_data(self):
        self.products[0]['material_url']='C:\\image.jpg'
        with self.assertRaises(ValueError): build_rows(self.config,self.products)
        self.products[0]['material_url']='https://example.com/3.jpg'
        self.config['variants'][0]['weight_g']=None
        with self.assertRaises(ValueError): build_rows(self.config,self.products)

    def test_excel_roundtrip(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as d:
            template=Path(d)/'template.xlsx'; output=Path(d)/'out.xlsx'
            rows=build_rows(self.config,self.products)
            w=Workbook();w.active.title='导入模板';w.active.append(list(rows[0]));w.create_sheet('说明')['A1']='保留';w.save(template)
            self.assertEqual(export(template,self.config,self.products,output),2)
            result=load_workbook(output)
            self.assertEqual(result['导入模板'].max_row,3)
            self.assertEqual(result['说明']['A1'].value,'保留')
            self.assertEqual(result['导入模板']['D2'].value,'001')
