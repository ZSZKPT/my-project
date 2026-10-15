import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from import_tools.miaoshou import export

class MiaoshouTest(unittest.TestCase):
    def test_export_links_defaults_weights_and_price(self):
        rows=[]
        for size,weight in [('S',570),('M',600)]:
            rows.append({'产品ID':'100','产品名称':'测试套装','主货号':'ABC','产品规格':'混合色|'+size,'供货价':'69CNY','尺寸':'长：33,宽：25, 高：2','毛重(g)':str(weight),'净重(g)':str(weight-10),'平台SKU':'ABC-'+size,'产品颜色图链接':'https://example.com/first.jpg\nhttps://example.com/second.jpg','产品素材图链接':'https://example.com/material.jpg','包装信息':'外包装形状：不规则\n外包装类型：软包装+软物\n外包装图片链接：https://example.com/package.jpg'})
        template=Path(__file__).resolve().parents[1]/'desktop/assets/妙手Temu服饰模板.xlsx'
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'out.xlsx';result=export(template,rows,out,{'测试套装':'Test set'})
            self.assertEqual(result['sku_count'],2)
            w=load_workbook(out);s=w['Sheet1'];h={str(c.value).strip():c.column for c in s[1] if c.value}
            for index in (3,4):
                self.assertEqual(s.cell(index,h['*产地']).value,'中国-江西省')
                self.assertEqual(s.cell(index,h['*定制品']).value,'否')
                self.assertEqual(s.cell(index,h['* 是否敏感属性']).value,'否')
                self.assertEqual(s.cell(index,h['* 颜色图']).value,rows[index-3]['产品颜色图链接'])
                self.assertEqual(s.cell(index,h['外包装图片']).value,'https://example.com/package.jpg')
                self.assertEqual(s.cell(index,h['* 供货价\n（CNY）']).value,69)
                self.assertEqual(s.cell(index,h['建议售价\n（（CNY））']).value,345)
            self.assertEqual(s.cell(4,h['* 毛重（g）']).value,600)
            self.assertIsNotNone(s['A2'].value)
            self.assertIsNone(s['D5'].value)
            with self.assertRaises(ValueError):export(template,rows,out,{'测试套装':'Test set'})
            out2=Path(d)/'override.xlsx';export(template,rows,out2,{'测试套装':'Test set'},price_override='88.5')
            self.assertEqual(load_workbook(out2)['Sheet1'].cell(3,h['* 供货价\n（CNY）']).value,88.5)
            self.assertEqual(load_workbook(out2)['Sheet1'].cell(3,h['建议售价\n（（CNY））']).value,442.5)
            with self.assertRaises(ValueError):export(template,rows,Path(d)/'missing.xlsx',{})
