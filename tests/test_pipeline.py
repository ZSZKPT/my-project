import json
import queue
import threading
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from openpyxl import load_workbook
from desktop.core import discover
from desktop.profiles import load,save,products_to_rows
from desktop.legacy import App

ASSETS=Path(__file__).resolve().parents[1]/'desktop/assets'

class PipelineTest(unittest.TestCase):
    def test_multi_template_saves_without_credentials(self):
        profile=load(ASSETS/'运动短袖套装.json')
        with tempfile.TemporaryDirectory() as d:
            first=save(d,profile)
            second_profile=dict(profile,name='运动套装二',category_id='29046',main_suffix='主图',key='not-a-real-key')
            second=save(d,second_profile)
            self.assertNotEqual(first,second)
            self.assertEqual(load(first)['name'],'运动短袖套装')
            self.assertEqual(load(second)['name'],'运动套装二')
            self.assertEqual(load(second)['main_suffix'],'主图')
            self.assertNotIn('not-a-real-key',second.read_text(encoding='utf-8'))
            self.assertEqual(len(list(Path(d).glob('*.json'))),2)

    def test_image_upload_to_miaoshou_workbook(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);source=root/'images'
            for sku in ('ABC','DEF'):
                folder=source/sku;folder.mkdir(parents=True)
                for name in [sku+'-前.jpg']+[sku+f'-图{i}.jpg' for i in range(1,5)]:Image.new('RGB',(1200,1600),'red').save(folder/name)
            groups=discover(source)
            for group,name in zip(groups,['字母印花','字母印花']):group['print_name']=name
            profile=load(ASSETS/'运动短袖套装.json');profile['price']=29
            app=App.__new__(App);app.events=queue.Queue();app.stop_retry=threading.Event()
            values={'output':str(root/'output'),'size':'1000','mode':'ImgBB 自动上传','links':'','key':'TEST-NOT-A-CREDENTIAL','template':str(ASSETS/'妙手Temu服饰模板.xlsx')}
            def upload(path,key):return 'https://example.com/'+Path(path).parent.name+'/'+Path(path).name
            with patch('desktop.legacy.upload_imgbb',side_effect=upload) as mocked:
                app.work(values,profile,groups)
                self.assertEqual(mocked.call_count,12)
            output=next((root/'output').iterdir());file=output/'妙手导入.xlsx'
            self.assertTrue(file.is_file(),list(app.events.queue))
            workbook=load_workbook(file);sheet=workbook['Sheet1'];headers={str(c.value).strip():c.column for c in sheet[1] if c.value}
            rows=[dict(zip([str(c.value).strip() for c in sheet[1]],row)) for row in sheet.iter_rows(min_row=3,values_only=True) if row[3]]
            self.assertEqual(len(rows),10)
            self.assertEqual(rows[0]['*产品标题'],'女士圆领短袖2件套字母印花圆领T恤+短裤')
            self.assertEqual(rows[-1]['平台SKU'],'DEF-XXL')
            self.assertEqual(rows[0]['*产品标题'],rows[-1]['*产品标题'])
            self.assertEqual({r['* 主编号'] for r in rows},{'ABC','DEF'})
            self.assertTrue(rows[0]['*英文标题'].endswith('ABC'))
            self.assertTrue(rows[-1]['*英文标题'].endswith('DEF'))
            self.assertNotEqual(rows[0]['* 颜色图'],rows[-1]['* 颜色图'])
            for row in rows:
                self.assertEqual(row['类目ID'],'29046')
                self.assertEqual(row['* 供货价\n（CNY）'],29)
                self.assertEqual(row['建议售价\n（（CNY））'],145)
                self.assertIn('素材图_1000x1000.jpg',row['* 产品素材图'])
                self.assertEqual(len(row['* 颜色图'].splitlines()),5)
                self.assertEqual(row['*产地'],'中国-江西省')
            self.assertNotIn('TEST-NOT-A-CREDENTIAL',(output/'模板配置.json').read_text(encoding='utf-8'))
            self.assertTrue(any('待手动导入' in str(message) for message in app.events.queue))
