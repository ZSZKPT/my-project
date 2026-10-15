import csv
from pathlib import Path


def read_names(path):
    path=Path(path)
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        workbook=load_workbook(path,read_only=True,data_only=True)
        try:rows=list(workbook.active.iter_rows(values_only=True))
        finally:workbook.close()
    elif path.suffix.lower()=='.csv':
        raw=path.read_bytes()
        for encoding in ('utf-8-sig','gb18030','utf-16'):
            try:rows=list(csv.reader(raw.decode(encoding).splitlines()));break
            except UnicodeError:continue
        else:raise ValueError('无法识别CSV编码')
    else:raise ValueError('请选择xlsx或csv文件')
    if not rows:raise ValueError('印花表为空')
    header=[str(v or '').strip() for v in rows[0]]
    sku_col=next((header.index(k) for k in ('货号','主货号','sku','SKC货号') if k in header),None)
    name_col=next((header.index(k) for k in ('印花名称','印花名','print_name') if k in header),None)
    if sku_col is None or name_col is None:raise ValueError('首行必须有“货号”和“印花名称”两列')
    result={}
    for index,row in enumerate(rows[1:],2):
        if not any(v not in (None,'') for v in row):continue
        sku=str(row[sku_col] or '').strip() if len(row)>sku_col else ''
        name=str(row[name_col] or '').strip() if len(row)>name_col else ''
        if not sku or not name:raise ValueError(f'第{index}行货号或印花名称为空')
        if sku in result:raise ValueError(f'第{index}行货号重复：{sku}')
        result[sku]=name
    if not result:raise ValueError('没有印花记录')
    return result


def title_for(rule, name, sku):
    if '{print_name}' not in rule:raise ValueError('标题规则必须包含 {print_name}')
    if rule.count('{print_name}')!=1:raise ValueError('标题规则只能包含一次 {print_name}')
    # Accept both 字母 and 字母印花 without producing 印花印花.
    if '{print_name}印花' in rule and name.endswith('印花'):name=name[:-2]
    return rule.replace('{print_name}',name).replace('{sku}',sku)
