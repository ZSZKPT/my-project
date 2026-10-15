import base64
import json
import re
import time
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
import ssl


class RetryableUploadError(ValueError):
    """A temporary upload failure that may safely resume from the checkpoint."""

from urllib.parse import urlencode
from PIL import Image, ImageOps
from import_tools.miaoshou import number

IMAGE_TYPES = {'.jpg', '.jpeg', '.png', '.webp'}


def natural_key(name):
    return [int(p) if p.isdigit() else p.casefold() for p in re.split(r'(\d+)', name)]


def discover(source, main_suffix="前"):
    main_suffix = main_suffix.strip().casefold()
    source = Path(source).resolve()
    if not source.is_dir():
        raise ValueError('请选择图片目录')
    folders = [source] if any(p.is_file() and p.suffix.lower() in IMAGE_TYPES and '素材图' not in p.stem for p in source.iterdir()) else sorted((p for p in source.iterdir() if p.is_dir()), key=lambda p: natural_key(p.name))
    groups = []
    for folder in folders:
        files = [p for p in folder.iterdir() if p.is_file() and not p.is_symlink() and p.suffix.lower() in IMAGE_TYPES and '素材图' not in p.stem]
        files.sort(key=lambda p: (not (main_suffix and p.stem.casefold().endswith(main_suffix)), natural_key(p.stem)))
        if files:
            groups.append({'sku':folder.name, 'files':files})
    if not groups:
        raise ValueError('未找到 JPG、PNG 或 WebP 图片；请选择货号目录或多个货号的父目录')
    return groups


def prepare(group, destination, size=1000):
    out = Path(destination) / group['sku']
    out.mkdir(parents=True, exist_ok=False)
    paths = []
    for index, source in enumerate(group['files'], 1):
        with Image.open(source) as original:
            original.load()
            image = ImageOps.exif_transpose(original).convert('RGB')
            path = out / f'{index:02d}_{group["sku"]}_{"主图" if index==1 else "商品图"}.jpg'
            image.save(path, quality=95)
            paths.append(path)
    with Image.open(paths[0]) as main:
        material = ImageOps.fit(main, (size, size), method=Image.Resampling.LANCZOS, centering=(.5,.5))
        material_path = out / f'素材图_{size}x{size}.jpg'
        material.save(material_path, quality=95)
    return {'sku': group['sku'], 'carousel':paths, 'material':material_path}


def upload_imgbb(path, key):
    if not key.strip():
        raise ValueError('自动上传需要 ImgBB API Key')
    # Documented API: https://api.imgbb.com/ . Never logs keys or response errors.
    payload = urlencode({'key':key.strip(),'image':base64.b64encode(Path(path).read_bytes()).decode(),'name':Path(path).stem}).encode()
    req = Request('https://api.imgbb.com/1/upload', data=payload, headers={'Content-Type':'application/x-www-form-urlencoded'})
    try:
        with urlopen(req, timeout=90) as response:
            result=json.load(response)
        url=result.get('data',{}).get('url','')
        if not result.get('success') or not url.startswith('https://'):
            raise ValueError('无有效图片链接')
        return url
    except HTTPError as error:
        if error.code in (408, 429) or 500 <= error.code <= 599:
            raise RetryableUploadError('图片服务暂时不可用') from None
        raise ValueError('图片服务拒绝上传，请检查 API Key、图片或额度') from None
    except URLError as error:
        if isinstance(error.reason, ssl.SSLCertVerificationError):
            raise ValueError('图片服务证书验证失败，请检查系统时间或证书配置') from None
        raise RetryableUploadError('图片上传网络暂时不可用') from None
    except (TimeoutError, ConnectionError):
        raise RetryableUploadError('图片上传连接中断或超时') from None
    except (ValueError, TypeError):
        raise ValueError('图片服务未返回有效链接，请检查 API Key、图片或额度') from None


def load_profile(path):
    from openpyxl import load_workbook
    sheet=load_workbook(path,data_only=True).active
    headers=[c.value for c in sheet[1]]
    required={'产品标题','变种属性值二','长','宽','高','重量'}
    if not required.issubset(headers):
        raise ValueError('此文件不是店小秘商品导出表')
    rows=[dict(zip(headers,row)) for row in sheet.iter_rows(min_row=2,values_only=True) if any(v is not None for v in row)]
    if not rows:
        raise ValueError('商品模板为空')
    if len({r['产品标题'] for r in rows}) != 1:
        raise ValueError('模板必须只含一个商品')
    sizes=[r['变种属性值二'] for r in rows]
    if len(set(sizes))!=len(sizes):
        raise ValueError('模板中有重复尺码，请只保留一个颜色的商品')
    profile = {'title':'女士圆领短袖2件套{print_name}印花圆领T恤+短裤','description':rows[0].get('产品描述') or '', 'color':rows[0].get('变种属性值一') or '混合色', 'origin':rows[0].get('产地') or '', 'variants':[{'size':r['变种属性值二'],'length_cm':number(r['长'],str(r['变种属性值二'])+'尺码的包装长'),'width_cm':number(r['宽'],str(r['变种属性值二'])+'尺码的包装宽'),'height_cm':number(r['高'],str(r['变种属性值二'])+'尺码的包装高'),'weight_g':number(r['重量'],str(r['变种属性值二'])+'尺码的毛重')} for r in rows]}
    # Carry only optional fields supported by the creation import format.
    optional = ['外包装形状','外包装类型','外包装图片','建议售价（USD）','净重（g）','是否定制品','定制品类型','一级工艺','二级工艺','产品视频url','描述视频url','产品说明书','说明书语种','SKU分类类型','SKU分类数量','SKU分类单位','是否独立包装','单品净含量','单品净含量单位','内计共含件数','是否同品','总净含量','总净含量单位','是否敏感属性','敏感属性值','储电容量','刀具长度','刀刃角度','液体容量','来源URL']
    profile['extra_fields'] = {key: rows[0][key] for key in optional if rows[0].get(key) not in (None, '')}
    return profile
