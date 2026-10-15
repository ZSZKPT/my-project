import queue
import sys
import threading
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
if __package__ in (None,''):sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from import_tools.miaoshou import read_csv, export, DEFAULT_ENGLISH
from desktop.print_names import read_names, title_for
ASSETS=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parent))/'assets'

class App:
    def __init__(self,window):
        self.window=window;window.title('飞凡上架 · Temu服饰');window.geometry('1100x760')
        self.rows=[];self.names={};self.events=queue.Queue();self.busy=False
        self.values={k:tk.StringVar(value=v) for k,v in {'source':'','output':'','template':str(ASSETS/'妙手Temu服饰模板.xlsx'),'origin':'中国-江西省','price':'','title_rule':'女士圆领短袖2件套{print_name}印花圆领T恤+短裤'}.items()}
        frame=ttk.Frame(window,padding=18);frame.pack(fill='both',expand=True)
        ttk.Label(frame,text='飞凡上架',font=('',18,'bold')).pack(anchor='w')
        ttk.Label(frame,text='读取SKU导出CSV → 核对英文标题与价格 → 生成妙手服饰导入表').pack(anchor='w',pady=8)
        for key,label,folder in [('source','SKU导出文件（CSV）',False),('output','保存目录',True),('template','妙手官方服饰模板',False)]:
            row=ttk.Frame(frame);row.pack(fill='x',pady=6);ttk.Label(row,text=label,width=25).pack(side='left')
            ttk.Entry(row,textvariable=self.values[key]).pack(side='left',fill='x',expand=True)
            ttk.Button(row,text='选择',command=lambda k=key,f=folder:self.choose(k,f)).pack(side='left')
        for key,label in [('origin','产地'),('price','统一供货价CNY（可留空）')]:
            row=ttk.Frame(frame);row.pack(fill='x',pady=6);ttk.Label(row,text=label,width=25).pack(side='left');ttk.Entry(row,textvariable=self.values[key]).pack(side='left',fill='x',expand=True)
        ttk.Label(frame,text='留空供货价：保留源表各SKU价格。建议售价自动为供货价5倍；非定制品、非敏感属性。').pack(anchor='w',pady=5)
        ttk.Label(frame,text='素材图、颜色图、外包装图片均取源CSV链接，不重传、不替换。颜色图为空时使用源轮播图/预览图。').pack(anchor='w')
        self.scan_button=ttk.Button(frame,text='读取SKU并核对标题',command=self.scan);self.scan_button.pack(anchor='w',pady=8)
        row=ttk.Frame(frame);row.pack(fill='x',pady=4)
        ttk.Label(row,text='印花标题规则',width=25).pack(side='left')
        ttk.Entry(row,textvariable=self.values['title_rule']).pack(side='left',fill='x',expand=True)
        ttk.Button(row,text='导入印花表格',command=self.import_names).pack(side='left')
        self.canvas=tk.Canvas(frame,height=180);self.canvas.pack(fill='both',expand=True)
        self.table=ttk.Frame(self.canvas);self.canvas.create_window((0,0),window=self.table,anchor='nw')
        scrollbar=ttk.Scrollbar(frame,orient='vertical',command=self.canvas.yview);scrollbar.pack(side='right',fill='y');self.canvas.configure(yscrollcommand=scrollbar.set)
        self.table.bind('<Configure>',lambda e:self.canvas.configure(scrollregion=self.canvas.bbox('all')))
        self.run=ttk.Button(frame,text='生成妙手上传表格',command=self.start);self.run.pack(anchor='w',pady=10)
        self.log=tk.Text(frame,height=9,state='disabled');self.log.pack(fill='x')
        ttk.Label(frame,text='本工具生成批量导入文件；不会登录妙手、提交发布或补造尺码测量表及类目属性。').pack(anchor='w',pady=5)
        ttk.Button(frame,text='图片自动上传 / 妙手多模板工具',command=self.local).pack(anchor='w')
        window.after(100,self.poll)

    def local(self):
        from desktop.legacy import App as LegacyApp
        child=tk.Toplevel(self.window);LegacyApp(child)

    def choose(self,key,folder):
        if self.busy:return
        value=filedialog.askdirectory() if folder else filedialog.askopenfilename(filetypes=[('CSV','*.csv')] if key=='source' else [('Excel','*.xlsx')])
        if value:self.values[key].set(value)
        if key=='source' and value:self.rows=[]

    def write(self,text):
        self.log.configure(state='normal');self.log.insert('end',text+'\n');self.log.see('end');self.log.configure(state='disabled')

    def scan(self):
        if self.busy:return
        try:
            rows=read_csv(self.values['source'].get());old={k:v.get() for k,v in self.names.items()};self.rows=rows;self.loaded_source=self.values['source'].get();self.names={}
            self.render_titles(old)
            self.write(f'读取 {len(rows)} 个SKU、{len(self.names)} 个标题。请填写或核对右侧英文标题。')
        except Exception as error:messagebox.showerror('读取失败',str(error))

    def render_titles(self,old):
        for widget in self.table.winfo_children():widget.destroy()
        for index,title in enumerate(dict.fromkeys(r['产品名称'] for r in self.rows)):
            ttk.Label(self.table,text=title,wraplength=370,width=45).grid(row=index,column=0,sticky='w',pady=6)
            var=tk.StringVar(value=old.get(title,DEFAULT_ENGLISH.get(title,'')));self.names[title]=var
            ttk.Entry(self.table,textvariable=var,width=75).grid(row=index,column=1,pady=6)


    def import_names(self):
        if self.busy:return
        try:
            if not self.rows:raise ValueError('请先读取SKU文件')
            path=filedialog.askopenfilename(filetypes=[('印花表格','*.xlsx *.csv')])
            if not path:return
            mapping=read_names(path);rule=self.values['title_rule'].get()
            matched={r['主货号'] for r in self.rows if r['主货号'] in mapping}
            if not matched:raise ValueError('没有匹配的货号')
            updates={sku:title_for(rule,mapping[sku],sku) for sku in matched}
            if not messagebox.askyesno('确认替换标题',f'按标题规则替换 {len(matched)} 个货号的中文标题，英文标题需重新核对。是否继续？'):return
            old={k:v.get() for k,v in self.names.items()}
            for r in self.rows:
                if r['主货号'] in updates:r['产品名称']=updates[r['主货号']]
            self.names={};self.render_titles(old)
            missing={r['主货号'] for r in self.rows}-matched
            extra=set(mapping)-matched
            self.write(f'已匹配{len(matched)}个货号；未匹配源货号：{sorted(missing)}；表格多余货号：{sorted(extra)}。未匹配商品保留原标题，英文标题请核对。')
        except Exception as error:messagebox.showerror('印花导入失败',str(error))

    def start(self):
        if self.busy:return
        try:
            if not self.rows or self.values['source'].get()!=self.loaded_source:raise ValueError('请先读取当前CSV')
            titles={k:v.get().strip() for k,v in self.names.items()}
            if not all(titles.values()):raise ValueError('请补齐所有英文标题')
            if not self.values['output'].get():raise ValueError('请选择保存目录')
            values={k:v.get().strip() for k,v in self.values.items()}
            self.busy=True;self.run.configure(state='disabled');self.scan_button.configure(state='disabled')
            threading.Thread(target=self.work,args=(values,list(self.rows),titles),daemon=True).start()
        except Exception as error:messagebox.showerror('请核对',str(error))

    def work(self,values,rows,titles):
        try:
            out=Path(values['output'])/('妙手Temu导入_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f')+'.xlsx')
            result=export(values['template'],rows,out,titles,values['origin'],values['price'] or None)
            self.events.put(f'已生成 {result["product_count"]} 个商品、{result["sku_count"]} 个SKU：{out}')
            for warning in result['warnings']:self.events.put('提示：'+warning)
            self.events.put('请在妙手导入核对；保留原模板说明和查找表，已移除示例商品。')
        except Exception as error:self.events.put('生成失败：'+str(error))
        finally:self.events.put(None)

    def poll(self):
        try:
            while True:
                event=self.events.get_nowait()
                if event is None:self.busy=False;self.run.configure(state='normal');self.scan_button.configure(state='normal')
                else:self.write(event)
        except queue.Empty:pass
        self.window.after(100,self.poll)

if __name__=='__main__':
    root=tk.Tk()
    from desktop.legacy import App as ImageApp
    image_app=ImageApp(root)
    if '--self-test' in sys.argv:
        root.withdraw();root.update()
        from openpyxl import load_workbook
        book=load_workbook(ASSETS/'妙手Temu服饰模板.xlsx',read_only=True)
        assert 'Sheet1' in book.sheetnames
        book.close()
        import tempfile
        from PIL import Image
        from desktop.thumbnails import open_preview
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)
            for index in range(12):
                folder=source/f'TEST-{index:02d}';folder.mkdir()
                for image_index in range(5):Image.new('RGB',(600,800),'#087f8c').save(folder/('前.jpg' if image_index==0 else f'图{image_index}.jpg'))
            image_app.variables['source'].set(str(source));image_app.scan()
            root.deiconify();root.geometry('1100x760');root.update()
            assert len(image_app.thumbnails)==12
            assert len(image_app.names)==12
            image_app.names[0].set('缩略图测试')
            before=image_app.canvas.yview()[0]
            entry=next(widget for widget in image_app.table.winfo_children() if isinstance(widget,ttk.Entry))
            entry.event_generate('<MouseWheel>',delta=-120);root.update()
            assert image_app.canvas.yview()[0]>before, 'Mouse wheel did not scroll print list'
            preview=open_preview(root,image_app.groups[0]['files'][0],'TEST-00');root.update()
            assert preview is not None
            second=open_preview(root,image_app.groups[1]['files'][0],'TEST-01');root.update()
            assert second is not None
            assert not preview.winfo_exists(), 'Previous preview was not closed'
            assert root._feifan_preview is second
            second.destroy()
        root.destroy()
    else:root.mainloop()
