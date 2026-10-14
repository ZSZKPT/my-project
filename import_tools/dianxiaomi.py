"""Generate Dianxiaomi import workbook using a user-supplied official template."""
import argparse
import csv
import json
import math
from pathlib import Path
from urllib.parse import urlsplit


def image_url(value):
    value = str(value or '').strip()
    parsed = urlsplit(value)
    if parsed.scheme not in ('https', 'http') or not parsed.hostname or parsed.username or parsed.password or parsed.hostname in ('localhost', '127.0.0.1'):
        raise ValueError('图片必须填写外网可访问的 HTTP(S) 链接，不支持本地路径')
    return value


def build_rows(config, products):
    title_rule = config.get('title', '')
    if '{print_name}' not in title_rule:
        raise ValueError('标题必须包含 {print_name}')
    variants = config.get('variants', [])
    if not variants:
        raise ValueError('至少需要一个尺码')
    rows, seen, titles = [], set(), set()
    for product in products:
        sku = str(product.get('sku', '')).strip()
        name = str(product.get('print_name', '')).strip()
        if not sku or not name or sku in seen:
            raise ValueError('货号与印花名称不能为空，货号不得重复')
        seen.add(sku)
        title = title_rule.replace('{print_name}', name).replace('{sku}', sku)
        if title in titles:
            raise ValueError('不同货号生成了相同标题，可能被合并；请在标题规则中增加 {sku} 或使用不同印花名')
        titles.add(title)
        carousel = [image_url(url) for url in str(product.get('carousel_urls', '')).split('|') if url.strip()]
        if not 1 <= len(carousel) <= 10:
            raise ValueError(f'{sku}: 轮播图需 1 至 10 个链接，用 | 分隔')
        material = image_url(product.get('material_url'))
        sizes = set()
        for variant in variants:
            size = str(variant.get('size', '')).strip()
            if not size or size in sizes:
                raise ValueError('尺码为空或重复')
            sizes.add(size)
            values = {}
            for key in ('price', 'length_cm', 'width_cm', 'height_cm', 'weight_g'):
                try:
                    value = float(variant.get(key, config.get(key, 0)))
                except (TypeError, ValueError):
                    raise ValueError(f'{sku}/{size}: {key} 尚未填写有效数值')
                if not math.isfinite(value) or value <= 0:
                    raise ValueError(f'{sku}/{size}: {key} 必须为正数')
                values[key] = value
            rows.append({
                '*产品标题': title, '英文标题': config.get('english_title', ''),
                '产品描述': config.get('description', ''), '产品货号': sku,
                '*变种属性名称一': '颜色', '*变种属性值一': config.get('color', '混合色'),
                '变种属性名称二': '尺寸', '变种属性值二': size,
                '预览图': image_url(product.get('preview_url') or carousel[0]),
                '*申报价格\n(店铺币种)': values['price'], 'SKU货号': f'{sku}-{size}',
                '*长（cm）': values['length_cm'], '*宽（cm）': values['width_cm'], '*高（cm）': values['height_cm'],
                '*重量（g）': values['weight_g'], '*轮播图': '\n'.join(carousel), '*产品素材图': material,
                '产地': config.get('origin', '')
            })
    if not rows:
        raise ValueError('商品表为空')
    return rows


def export(template, config, products, output):
    from openpyxl import load_workbook
    rows = build_rows(config, products)
    workbook = load_workbook(template)
    sheet = workbook['导入模板']
    headers = [cell.value for cell in sheet[1]]
    missing = set(rows[0]) - set(headers)
    if missing:
        raise ValueError('模板缺少字段: ' + ', '.join(sorted(missing)))
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append([row.get(header, '') for header in headers])
        # Force text: names and SKUs beginning with '=' must not become formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    workbook.save(output)
    return len(rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', required=True)
    parser.add_argument('--config', required=True)
    parser.add_argument('--products', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    output = Path(args.output)
    if output.exists():
        parser.error('输出文件已存在，请换一个文件名，避免覆盖')
    if output.resolve() in {Path(p).resolve() for p in (args.template, args.config, args.products)}:
        parser.error('输出路径不能覆盖输入文件')
    config = json.loads(Path(args.config).read_text(encoding='utf-8-sig'))
    with open(args.products, encoding='utf-8-sig', newline='') as file:
        products = list(csv.DictReader(file))
    try:
        count = export(args.template, config, products, output)
    except (ValueError, KeyError) as error:
        parser.error(str(error))
    print(f'已生成 {count} 行 SKU。图片可访问性、素材图尺寸和店小秘实际导入尚需核对；未发布商品。')

if __name__ == '__main__':
    main()
