"""Map exported Temu SKU rows to the official Miaoshou apparel workbook."""
import csv
import math
import re
import unicodedata
from copy import copy
from decimal import Decimal, ROUND_DOWN
from pathlib import Path
from openpyxl import load_workbook

REQUIRED = {'产品ID','产品名称','主货号','产品规格','供货价','尺寸','毛重(g)','净重(g)','平台SKU','包装信息'}
DEFAULT_ENGLISH = {'女士圆领卫衣套装色块爱心花朵案印花长袖卫衣+直筒长裤': 'Women’s Crew Neck Sweatshirt and Straight-Leg Pants Set with Color-Block Heart and Floral Print'}


def read_csv(path):
    raw=Path(path).read_bytes()
    for encoding in ('utf-8-sig','gb18030','utf-16'):
        try:
            text=raw.decode(encoding)
            reader=csv.DictReader(text.splitlines(keepends=True))
            if not REQUIRED.issubset(reader.fieldnames or []):continue
            rows=[{k:(v or '').strip() for k,v in r.items() if k} for r in reader]
            if not rows:raise ValueError('SKU 表为空')
            return rows
        except UnicodeError:continue
    raise ValueError('请选择含产品ID、产品名称、产品规格、供货价等字段的SKU导出CSV')


def number(value, label, positive=True):
    if value is None or (isinstance(value,str) and not value.strip()):raise ValueError(f'{label}不能为空')
    if isinstance(value,bool):raise ValueError(f'{label}不是有效数字')
    if isinstance(value,str):value=unicodedata.normalize('NFKC',value).strip()
    try:value=float(value)
    except (ValueError,TypeError):raise ValueError(f'{label} 不是有效数字')
    if not math.isfinite(value) or (value<=0 if positive else value<0):raise ValueError(f'{label} 数值无效')
    return value


def parse_price(value,label='统一申报价'):
    if value is None or not str(value).strip():raise ValueError(f'{label}不能为空，请填写例如29或29.90')
    if isinstance(value,str):
        text=unicodedata.normalize('NFKC',value).strip()
        match=re.fullmatch(r'(?:[¥￥]\s*)?(\d+(?:\.\d+)?)\s*(?:元|CNY)?',text,re.IGNORECASE)
        if not match:raise ValueError(f'{label}格式不正确，请填写例如29或29.90，不要使用逗号')
        value=match.group(1)
    return number(value,label)


def export(template, rows, output, english_titles, origin='中国-江西省', price_override=None):
    if not rows:raise ValueError('没有SKU数据')
    if Path(output).exists():raise ValueError('输出文件已存在，请换文件名')
    if origin!='中国-江西省' and not origin.strip():raise ValueError('产地不能为空')
    if price_override is not None:price_override=parse_price(price_override,'统一供货价')
    w=load_workbook(template);s=w['Sheet1']
    headers={str(c.value).strip():c.column for c in s[1] if c.value}
    categories={str(name):str(cid) for cid,name in w['类目ID'].iter_rows(min_row=2,values_only=True) if name}
    category_ids=set(categories.values())
    needed=['* 主编号','*产品标题','*英文标题','*产地','*定制品','* 是否敏感属性','平台SKU','* 颜色图']
    if any(h not in headers for h in needed):raise ValueError('请选择妙手Temu服饰类官方模板')
    styles={c.column:copy(c._style) for c in s[3]}
    for merged in list(s.merged_cells.ranges):
        if merged.min_row>=3:s.unmerge_cells(str(merged))
    for row in s.iter_rows(min_row=3):
        for c in row:
            c.hyperlink=None
            c.value=None
    seen=set();warnings=[]
    for index,r in enumerate(rows,3):
        title=r.get('产品名称','').strip();english=(r.get('英文标题') or english_titles.get(title,'')).strip()
        if not english:raise ValueError(f'请填写英文标题：{title}')
        specs=r.get('产品规格','').split('|')
        if len(specs)!=2 or not all(specs):raise ValueError('产品规格必须为 颜色|尺码')
        identity=(r.get('产品ID'),r.get('主货号'),*specs)
        if identity in seen:raise ValueError('发现重复的商品/颜色/尺码组合')
        seen.add(identity)
        dimensions=dict(re.findall(r'(长|宽|高)\s*[：:]\s*([\d.]+)',r.get('尺寸','')))
        if len(dimensions)!=3:raise ValueError('无法识别包装长宽高')
        package=dict(re.findall(r'(外包装形状|外包装类型|外包装图片链接)\s*[：:]\s*([^\n]+)',r.get('包装信息','')))
        match=re.fullmatch(r'(\d+(?:\.\d+)?)\s*CNY',r.get('供货价',''))
        if price_override is None and not match:raise ValueError('供货价需为CNY，请填写统一价格或修正源文件')
        category=str(r.get('类目ID') or '').strip()
        if category and category not in category_ids:raise ValueError('模板类目ID不在官方类目表中')
        category=category or categories.get(r.get('产品分类',''),'')
        if not category:warnings.append('部分商品类目路径未匹配，需在妙手选择类目')
        colors=r.get('产品颜色图链接','')
        if not colors:
            colors=r.get('产品轮播图链接','') or r.get('产品预览图链接','')
            if colors:warnings.append('部分颜色图为空，使用源文件轮播图/预览图回填')
        if not colors:raise ValueError('源文件缺少颜色图、轮播图和预览图')
        price = price_override if price_override is not None else number(match.group(1),'供货价')
        price = Decimal(str(price)).quantize(Decimal('0.01'), rounding=ROUND_DOWN)
        suggested = price * 5
        data={'类目ID':category,'* 主编号':r.get('产品ID') or r.get('主货号'),'*产品标题':title,'*英文标题':english,
        '产品描述':r.get('产品描述',''),'*产地':origin,'制造地点':r.get('制造地点',''),'站外产品链接':r.get('站外产品链接',''),
        '* 产品素材图':r.get('产品素材图链接',''),'*定制品':'否','* 规格名称1':'颜色','* 规格属性值1':specs[0],
        '* 规格名称2':'尺码','* 规格属性值2':specs[1],'* 颜色图':colors,'主货号':r.get('主货号',''),
        '* 供货价\n（CNY）':float(price),'建议售价\n（（CNY））':float(suggested),
        '* 长（cm）':number(dimensions['长'],'长'),'* 宽（cm）':number(dimensions['宽'],'宽'),'* 高（cm）':number(dimensions['高'],'高'),
        '* 毛重（g）':number(r.get('毛重(g)'),'毛重'),'净重(g)':number(r['净重(g)'],'净重',False) if r.get('净重(g)') else '',
        '平台SKU':r.get('平台SKU',''),'* 是否敏感属性':'否',
        '包装信息':package.get('外包装形状','').strip(),'外包装类型':package.get('外包装类型','').strip(),
        '外包装图片':package.get('外包装图片链接','').strip()}
        for key,value in data.items():
            cell=s.cell(index,headers[key],value);cell._style=copy(styles[cell.column])
            if isinstance(value,str):cell.data_type='s'
        s.row_dimensions[index].height=72
    w.save(output)
    return {'sku_count':len(rows),'product_count':len({r.get('产品ID') or r.get('主货号') for r in rows}),'warnings':list(dict.fromkeys(warnings))}
